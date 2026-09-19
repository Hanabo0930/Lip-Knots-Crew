"""依頼Excelの全シートを読取り、数式や欠落を正常な確定値にしない。"""
import argparse
import datetime as dt
import hashlib
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook


def extract_xlsx(filename):
    path = Path(filename)
    raw = path.read_bytes()
    if len(raw) > 25 * 1024 * 1024 or not raw.startswith(b"PK\x03\x04"):
        raise ValueError("Excelの実体またはサイズが不正です")
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(entry.file_size for entry in entries) > 100 * 1024 * 1024:
            raise ValueError("Excelの展開サイズが上限を超えています")
        if "xl/workbook.xml" not in archive.namelist():
            raise ValueError("Excelブックの構造を確認できません")
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    if not 1 <= len(book.worksheets) <= 100:
        raise ValueError("Excelのシート数が上限外です")
    sheets, issues, total = [], [], 0
    try:
        for sheet in book.worksheets:
            if sheet.max_row > 5000 or sheet.max_column > 100:
                raise ValueError("Excelの行列数が上限を超えています")
            rows = []
            for row_number, cells in enumerate(sheet.iter_rows(), 1):
                values = []
                for column_number, cell in enumerate(cells, 1):
                    total += 1
                    if total > 100000:
                        raise ValueError("Excelのセル数が上限を超えています")
                    value = cell.value
                    if cell.data_type == "f":
                        issues.append(f"{sheet.title}!{cell.coordinate} は数式です")
                        value = ""
                    elif cell.data_type == "e":
                        issues.append(f"{sheet.title}!{cell.coordinate} はエラー値です")
                        value = ""
                    elif isinstance(value, dt.datetime):
                        value = value.strftime("%Y-%m-%d %H:%M:%S")
                    elif isinstance(value, dt.date):
                        value = value.isoformat()
                    elif isinstance(value, dt.time):
                        value = value.strftime("%H:%M:%S")
                    elif value is None:
                        value = ""
                    elif not isinstance(value, (str, int, float, bool)):
                        issues.append(f"{sheet.title}!{cell.coordinate} の値を読めません")
                        value = ""
                    values.append(value)
                rows.append(values)
            sheets.append({"name": sheet.title, "state": sheet.sheet_state, "rows": rows})
    finally:
        book.close()
    return {"version": 1, "format": "xlsx", "sha256": hashlib.sha256(raw).hexdigest(),
            "bytes": len(raw), "complete": len(issues) == 0, "issues": issues, "sheets": sheets}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("input")
    parser.add_argument("output")
    args = parser.parse_args()
    result = extract_xlsx(args.input)
    Path(args.output).write_text(json.dumps(result, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"sheets": len(result["sheets"]), "complete": result["complete"]}, ensure_ascii=False))
